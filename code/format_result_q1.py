"""
问题 1 结果格式化：生成论文用表 1、表 2，并写入 result1.xlsx 模板。
"""
import numpy as np
import pandas as pd
from pathlib import Path
from openpyxl import load_workbook

from config import (
    SLOTS_PER_DAY,
    HOURS_PER_SLOT,
    RESULT_TEMPLATE_DIR,
    ROOT,
    PAPER_TABLE1_SLOTS,
    PAPER_TABLE2_RANGES,
)


def make_paper_table1(sol: dict) -> pd.DataFrame:
    """生成论文表 1：指定时段购电量 + 全天购电量/电费。"""
    E_grid = sol["E_grid_kWh"]
    price = sol["price"]
    rows = []
    for label, slot in PAPER_TABLE1_SLOTS.items():
        rows.append({"时间段": label, "购电量": E_grid[slot]})
    df = pd.DataFrame(rows)
    df["购电量"] = df["购电量"].round(2)

    total_kwh = round(E_grid.sum(), 2)
    total_cost = round((E_grid * price).sum(), 2)  # 与 solver 返回一致
    summary = pd.DataFrame([
        {"时间段": "全天购电量", "购电量": total_kwh},
        {"时间段": "全天购电费", "购电量": total_cost},
    ])
    return pd.concat([df, summary], ignore_index=True)


def make_paper_table2(sol: dict) -> pd.DataFrame:
    """生成论文表 2：指定时段充放电量 + 0:00/24:00 储电量。"""
    E_ch = sol["E_ch_kWh"]
    E_dis = sol["E_dis_kWh"]
    E = sol["E_kWh"]
    rows = []
    for label, start, end in PAPER_TABLE2_RANGES:
        rows.append({
            "时间段": label,
            "充电量": E_ch[start:end + 1].sum(),
            "放电量": E_dis[start:end + 1].sum(),
        })
    df = pd.DataFrame(rows)
    df["充电量"] = df["充电量"].round(2)
    df["放电量"] = df["放电量"].round(2)

    soc_rows = pd.DataFrame([
        {"时间段": "0:00 储电量", "充电量": round(E[0], 2), "放电量": np.nan},
        {"时间段": "24:00 储电量", "充电量": round(E[-1], 2), "放电量": np.nan},
    ])
    return pd.concat([df, soc_rows], ignore_index=True)


def write_result1_xlsx(sol: dict, out_path: Path = None, template_path: Path = None):
    """
    复制 result1.xlsx 模板并写入结果。
    工作表 1：计划购电量（144 行，每行对应 slot 0..143）
    工作表 2：充放电量（6 个时段累计 + 0:00/24:00 储电量）
    """
    if template_path is None:
        template_path = RESULT_TEMPLATE_DIR / "result1.xlsx"
    if out_path is None:
        out_path = ROOT / "result1.xlsx"

    # 复制模板
    import shutil
    shutil.copy(template_path, out_path)

    wb = load_workbook(out_path)

    # ---- 工作表 1：计划购电量 ----
    ws1 = wb.worksheets[0]
    E_grid = sol["E_grid_kWh"]
    # 模板第 1 行为表头，数据从第 2 行开始；按行顺序填入 slot 0..143
    for slot in range(SLOTS_PER_DAY):
        row = slot + 2
        ws1.cell(row=row, column=2, value=float(E_grid[slot]))

    # ---- 工作表 2：充放电量 ----
    ws2 = wb.worksheets[1]
    E_ch = sol["E_ch_kWh"]
    E_dis = sol["E_dis_kWh"]
    E = sol["E_kWh"]

    # 左半部分：6 个时段累计（第 2-7 行）
    for i, (label, start, end) in enumerate(PAPER_TABLE2_RANGES):
        row = i + 2
        ws2.cell(row=row, column=2, value=float(E_ch[start:end + 1].sum()))
        ws2.cell(row=row, column=3, value=float(E_dis[start:end + 1].sum()))

    # 右半部分：0:00 / 24:00 储电量（第 2-3 行）
    ws2.cell(row=2, column=5, value=float(E[0]))
    ws2.cell(row=3, column=5, value=float(E[-1]))

    wb.save(out_path)
    wb.close()
    return out_path


def format_all(sol: dict, out_dir: Path = None):
    """生成全部输出：表 1、表 2、result1.xlsx。"""
    if out_dir is None:
        out_dir = ROOT
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    table1 = make_paper_table1(sol)
    table2 = make_paper_table2(sol)
    result_path = write_result1_xlsx(sol, out_path=out_dir / "result1.xlsx")

    # 保存 CSV 备份（便于论文直接读取）
    table1.to_csv(out_dir / "table1_q1.csv", index=False)
    table2.to_csv(out_dir / "table2_q1.csv", index=False)

    return {
        "table1": table1,
        "table2": table2,
        "result_xlsx": result_path,
    }


def print_paper_tables(sol: dict):
    """在控制台打印表 1、表 2（供快速查看）。"""
    table1 = make_paper_table1(sol)
    table2 = make_paper_table2(sol)

    print("\n========== 表 1：微网在指定时间段的购电量及全天的购电量和购电费 ==========")
    print(table1.to_string(index=False))

    print("\n========== 表 2：储能设备在指定时间段的充放电量及 0:00 和 24:00 的储电量 ==========")
    print(table2.to_string(index=False))


if __name__ == "__main__":
    from solver_q1 import solve_problem1
    sol = solve_problem1()
    res = format_all(sol)
    print_paper_tables(sol)
    print(f"\n结果已保存到: {res['result_xlsx']}")
