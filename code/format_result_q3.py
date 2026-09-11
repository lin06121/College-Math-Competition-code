"""
问题 3 结果格式化：
- 论文表 1（指定时段购电量 + 全天购电量/购电费，4 个指定日期）
- 论文表 2（指定时段充放电量 + 0:00/24:00 储电量，4 个指定日期）
- 论文表 3（指定日期的紧急购电时间段与购电量）
- result3.xlsx（计划购电量 / 调整购电量 / 充放电量 / 紧急购电量 四个工作表）

结算口径（字面读法）：只按实际购入量 P_adj 付全价，偏差部分少买按 50%、多买按 150%：
    费用(t) = c·P_adj + 0.5·c·|P_adj − P_plan| + 5·c·P_em
表 1 的“全天购电费”为计划+调整口径（不含紧急购电，紧急另见表 3）。

模板列约定与 result2 一致（沿用已被接受的列映射）：
“计划购电量”/“调整购电量”工作表第 k 个时间列（k = 1..144）对应当天第 k-1 个
10 分钟时段，即第 1 列（模板标签 0:10-0:20）填入 0:00-0:10 时段，以此类推。
"""
import numpy as np
import pandas as pd
from pathlib import Path
import shutil
from openpyxl import load_workbook

from config import (
    SLOTS_PER_DAY, HOURS_PER_SLOT, RESULT_TEMPLATE_DIR, ROOT, KEY_DATES,
    PAPER_TABLE1_SLOTS, PAPER_TABLE2_RANGES,
)
from format_result_q2 import (
    EM_TOL_KW, slot_range_label, _date_label, _day_df, make_emergency_events,
)


# ----------------------------------------------------------------------------
# 结算费用（与 solver_q3.compute_costs 同口径）
# ----------------------------------------------------------------------------

def settled_fee(df: pd.DataFrame, include_emergency: bool = False):
    """全天（计划+调整）购电费与购电量；include_emergency 时并入紧急购电费。"""
    dt = HOURS_PER_SLOT
    P_plan = df["P_plan_kW"].values
    P_adj = df["P_adj_kW"].values
    c = df["price"].values
    fee = float((c * P_adj * dt).sum())
    fee += float((0.5 * c * np.abs(P_adj - P_plan) * dt).sum())
    if include_emergency:
        fee += float((5.0 * c * df["P_em_kW"].values * dt).sum())
    kwh = float(P_adj.sum() * dt)
    return kwh, fee


# ----------------------------------------------------------------------------
# 论文表 1 / 表 2 / 表 3
# ----------------------------------------------------------------------------

def make_paper_table1_q3(result_df, summary_df, key_dates=KEY_DATES) -> pd.DataFrame:
    """表 1：指定日期各指定时段的购电量（调整后）及全天购电量/购电费。"""
    rows = []
    for label, slot in PAPER_TABLE1_SLOTS.items():
        row = {"时间段": label}
        for d in key_dates:
            df = _day_df(result_df, d)
            row[_date_label(d)] = float(df.loc[slot, "P_adj_kW"] * HOURS_PER_SLOT)
        rows.append(row)

    tot_row = {"时间段": "全天购电量"}
    fee_row = {"时间段": "全天购电费"}
    for d in key_dates:
        df = _day_df(result_df, d)
        kwh, fee = settled_fee(df, include_emergency=False)
        tot_row[_date_label(d)] = kwh
        fee_row[_date_label(d)] = fee
    df_out = pd.DataFrame(rows + [tot_row, fee_row])
    return df_out[["时间段"] + [_date_label(d) for d in key_dates]].round(2)


def make_paper_table2_q3(result_df, summary_df, key_dates=KEY_DATES) -> pd.DataFrame:
    """表 2：指定日期各 4 小时时段充放电量及 0:00/24:00 储电量（实际执行）。"""
    rows = []
    for label, start, end in PAPER_TABLE2_RANGES:
        row = {"时间段": label}
        for d in key_dates:
            df = _day_df(result_df, d)
            row[f"{_date_label(d)}_充电量"] = float(
                (df.loc[start:end, "P_ch_kW"] * HOURS_PER_SLOT).sum())
            row[f"{_date_label(d)}_放电量"] = float(
                (df.loc[start:end, "P_dis_kW"] * HOURS_PER_SLOT).sum())
        rows.append(row)

    for soc_label, col in (("0:00 储电量", "e0"), ("24:00 储电量", "e_end")):
        row = {"时间段": soc_label}
        for d in key_dates:
            summ = summary_df[summary_df["date"] == pd.Timestamp(d)].iloc[0]
            row[f"{_date_label(d)}_充电量"] = float(summ[col])
            row[f"{_date_label(d)}_放电量"] = np.nan
        rows.append(row)

    cols = ["时间段"]
    for d in key_dates:
        cols += [f"{_date_label(d)}_充电量", f"{_date_label(d)}_放电量"]
    return pd.DataFrame(rows)[cols].round(2)


def make_paper_table3_q3(result_df, summary_df, key_dates=KEY_DATES,
                         tol: float = EM_TOL_KW) -> pd.DataFrame:
    """表 3：指定日期的紧急购电时间段、购电量与全天紧急购电费。"""
    ev = make_emergency_events(result_df, tol=tol)
    dt = HOURS_PER_SLOT
    rows = []
    for d in key_dates:
        df = _day_df(result_df, d)
        em_fee = float((5.0 * df["price"].values * df["P_em_kW"].values * dt).sum())
        sub = ev[ev["date"] == pd.Timestamp(d)]
        if sub.empty:
            rows.append({"日期": _date_label(d), "时间段": "无", "紧急购电量": 0.0,
                         "全天紧急购电费": round(em_fee, 2)})
        else:
            for i, (_, r) in enumerate(sub.iterrows()):
                rows.append({
                    "日期": _date_label(d) if i == 0 else "",
                    "时间段": r["时间段"],
                    "紧急购电量": round(float(r["紧急购电量"]), 2),
                    "全天紧急购电费": round(em_fee, 2) if i == 0 else "",
                })
    return pd.DataFrame(rows)


# ----------------------------------------------------------------------------
# result3.xlsx
# ----------------------------------------------------------------------------

def _write_qty_sheet(ws, dates, result_df, col, settle: bool):
    """写入 335×147 的“购电量”工作表（计划 or 调整）。"""
    for i, d in enumerate(dates):
        df = _day_df(result_df, d)
        row = i + 2
        ws.cell(row=row, column=1, value=pd.Timestamp(d).to_pydatetime())
        kwh = df[col].values * HOURS_PER_SLOT
        for t in range(SLOTS_PER_DAY):
            ws.cell(row=row, column=t + 2, value=float(kwh[t]))
        ws.cell(row=row, column=146, value=float(kwh.sum()))
        if settle:
            fee = settled_fee(df, include_emergency=False)[1]
        else:
            fee = float((df["P_plan_kW"].values * df["price"].values
                         * HOURS_PER_SLOT).sum())
        ws.cell(row=row, column=147, value=fee)


def write_result3_xlsx(result_df, summary_df, out_path: Path = None,
                       template_path: Path = None):
    """复制附件 5 模板并写入四个工作表。"""
    if template_path is None:
        template_path = RESULT_TEMPLATE_DIR / "result3.xlsx"
    if out_path is None:
        out_path = ROOT / "result3.xlsx"
    shutil.copy(template_path, out_path)
    wb = load_workbook(out_path)

    dates = sorted(result_df["date"].unique())

    # ---- 1：计划购电量（P_plan） ----
    _write_qty_sheet(wb["计划购电量"], dates, result_df, "P_plan_kW", settle=False)

    # ---- 2：调整购电量（P_adj，全天购电费为计划+调整结算费） ----
    _write_qty_sheet(wb["调整购电量"], dates, result_df, "P_adj_kW", settle=True)

    # ---- 3：充放电量（每天 6 个 4 小时时段 + 0:00/24:00 储电量） ----
    ws2 = wb["充放电量"]
    if ws2.max_row > 1:
        ws2.delete_rows(2, ws2.max_row - 1)
    r = 2
    for d in dates:
        df = _day_df(result_df, d)
        summ = summary_df[summary_df["date"] == pd.Timestamp(d)].iloc[0]
        for j, (label, start, end) in enumerate(PAPER_TABLE2_RANGES):
            ws2.cell(row=r, column=1,
                     value=pd.Timestamp(d).to_pydatetime() if j == 0 else None)
            ws2.cell(row=r, column=2, value=label)
            ws2.cell(row=r, column=3, value=float(
                (df.loc[start:end, "P_ch_kW"] * HOURS_PER_SLOT).sum()))
            ws2.cell(row=r, column=4, value=float(
                (df.loc[start:end, "P_dis_kW"] * HOURS_PER_SLOT).sum()))
            if j == 0:
                ws2.cell(row=r, column=5, value=pd.Timestamp(d).to_pydatetime())
                ws2.cell(row=r, column=6, value=float(summ["e0"]))
            if j == 1:
                ws2.cell(row=r, column=5, value="24:00")
                ws2.cell(row=r, column=6, value=float(summ["e_end"]))
            r += 1

    # ---- 4：紧急购电量 ----
    ws3 = wb["紧急购电量"]
    if ws3.max_row > 1:
        ws3.delete_rows(2, ws3.max_row - 1)
    ev = make_emergency_events(result_df)
    r = 2
    for _, e in ev.iterrows():
        ws3.cell(row=r, column=1, value=pd.Timestamp(e["date"]).to_pydatetime())
        ws3.cell(row=r, column=2, value=e["时间段"])
        ws3.cell(row=r, column=3, value=round(float(e["紧急购电量"]), 2))
        r += 1

    wb.save(out_path)
    wb.close()
    return out_path


# ----------------------------------------------------------------------------
# 一键输出
# ----------------------------------------------------------------------------

def format_all(result_df, summary_df, out_dir: Path = None):
    if out_dir is None:
        out_dir = ROOT
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    t1 = make_paper_table1_q3(result_df, summary_df)
    t2 = make_paper_table2_q3(result_df, summary_df)
    t3 = make_paper_table3_q3(result_df, summary_df)
    events = make_emergency_events(result_df)
    xlsx = write_result3_xlsx(result_df, summary_df, out_path=out_dir / "result3.xlsx")

    t1.to_csv(out_dir / "table1_q3.csv", index=False, encoding="utf-8-sig")
    t2.to_csv(out_dir / "table2_q3.csv", index=False, encoding="utf-8-sig")
    t3.to_csv(out_dir / "table3_q3.csv", index=False, encoding="utf-8-sig")
    events.to_csv(out_dir / "emergency_events_q3.csv", index=False,
                  encoding="utf-8-sig")
    return {"table1": t1, "table2": t2, "table3": t3,
            "events": events, "result_xlsx": xlsx}


def print_paper_tables(result_df, summary_df):
    print("\n===== 表 1：指定日期的购电量及全天购电量/购电费（kWh / 元） =====")
    print(make_paper_table1_q3(result_df, summary_df).to_string(index=False))
    print("\n===== 表 2：指定日期各时段充放电量及 0:00/24:00 储电量（kWh） =====")
    print(make_paper_table2_q3(result_df, summary_df).to_string(index=False))
    print("\n===== 表 3：指定日期的紧急购电量（kWh） =====")
    print(make_paper_table3_q3(result_df, summary_df).to_string(index=False))


if __name__ == "__main__":
    from config import DATA_PROCESSED
    rd = pd.read_csv(DATA_PROCESSED / "result3_details.csv", parse_dates=["date"])
    sd = pd.read_csv(DATA_PROCESSED / "result3_summary.csv", parse_dates=["date"])
    res = format_all(rd, sd)
    print_paper_tables(rd, sd)
    print(f"\n结果已保存到: {res['result_xlsx']}")
