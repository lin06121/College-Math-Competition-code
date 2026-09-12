"""
问题 2 结果格式化：
- 论文表 1（指定时段计划购电量 + 全天购电量/购电费，4 个指定日期）
- 论文表 2（指定时段充放电量 + 0:00/24:00 储电量，4 个指定日期）
- 论文表 3（指定日期的紧急购电时间段与购电量）
- result2.xlsx（计划购电量 / 充放电量 / 紧急购电量 三个工作表）

模板列约定（与附件 2 的“时段结束时刻”标记一致）：
“计划购电量”工作表第 k 个时间列（k = 1..144）对应当天第 k-1 个 10 分钟时段，
即第 1 列（标签 0:10-0:20）填入 0:00-0:10 时段，以此类推。
"""
import numpy as np
import pandas as pd
from pathlib import Path
import shutil
from openpyxl import load_workbook

from config import (
    SLOTS_PER_DAY,
    HOURS_PER_SLOT,
    RESULT_TEMPLATE_DIR,
    ROOT,
    KEY_DATES,
    PAPER_TABLE1_SLOTS,
    PAPER_TABLE2_RANGES,
)

# 判定“存在紧急购电”的功率阈值（kW），滤除数值噪声
EM_TOL_KW = 0.5


# ----------------------------------------------------------------------------
# 时间标签工具
# ----------------------------------------------------------------------------

def _hm(minutes: int) -> str:
    """把分钟数格式化为 H:MM。"""
    minutes = int(round(minutes))
    h, m = divmod(minutes, 60)
    return f"{h}:{m:02d}"


def slot_range_label(t0: int, t1: int) -> str:
    """时段 t0..t1（含，0-based）对应的区间标签，如 13:00-13:30。"""
    return f"{_hm(t0 * 10)}-{_hm((t1 + 1) * 10)}"


def _day_df(result_df: pd.DataFrame, date) -> pd.DataFrame:
    sub = result_df[result_df["date"] == pd.Timestamp(date)]
    return sub.sort_values("slot").reset_index(drop=True)


# ----------------------------------------------------------------------------
# 论文表 1 / 表 2
# ----------------------------------------------------------------------------

def make_paper_table1_q2(result_df, summary_df, key_dates=KEY_DATES) -> pd.DataFrame:
    """表 1：4 个指定日期在各指定时段的计划购电量及全天购电量/购电费（kWh / 元）。"""
    rows = []
    for label, slot in PAPER_TABLE1_SLOTS.items():
        row = {"时间段": label}
        for d in key_dates:
            df = _day_df(result_df, d)
            row[_date_label(d)] = float(df.loc[slot, "P_plan_kW"] * HOURS_PER_SLOT)
        rows.append(row)

    # 全天购电量 / 购电费
    tot_row = {"时间段": "全天购电量"}
    fee_row = {"时间段": "全天购电费"}
    for d in key_dates:
        df = _day_df(result_df, d)
        dt = HOURS_PER_SLOT
        tot_row[_date_label(d)] = float((df["P_plan_kW"] * dt).sum())
        fee_row[_date_label(d)] = float((df["P_plan_kW"] * df["price"] * dt).sum())
    df_out = pd.DataFrame(rows + [tot_row, fee_row])
    df_out = df_out[["时间段"] + [_date_label(d) for d in key_dates]]
    return df_out.round(2)


def _date_label(d) -> str:
    t = pd.Timestamp(d)
    return f"{t.year}.{t.month}.{t.day}"


def make_paper_table2_q2(result_df, summary_df, key_dates=KEY_DATES) -> pd.DataFrame:
    """表 2：4 个指定日期各 4 小时时段的充放电量及 0:00/24:00 储电量（kWh）。"""
    rows = []
    for label, start, end in PAPER_TABLE2_RANGES:
        row = {"时间段": label}
        for d in key_dates:
            df = _day_df(result_df, d)
            row[f"{_date_label(d)}_充电量"] = float(
                (df.loc[start:end, "P_ch_kW"] * HOURS_PER_SLOT).sum())
            row[f"{_date_label(d)}_放电量"] = float(
                (df.loc[start:end, "P_dis_kW"] * HOURS_PER_SLOT).sum())
        rows.append(row)

    for soc_label, col in (("0:00 储电量", "e0"), ("24:00 储电量", "e_end")):
        row = {"时间段": soc_label}
        for d in key_dates:
            summ = summary_df[summary_df["date"] == pd.Timestamp(d)].iloc[0]
            row[f"{_date_label(d)}_充电量"] = float(summ[col])
            row[f"{_date_label(d)}_放电量"] = np.nan
        rows.append(row)

    cols = ["时间段"]
    for d in key_dates:
        cols += [f"{_date_label(d)}_充电量", f"{_date_label(d)}_放电量"]
    df_out = pd.DataFrame(rows)[cols]
    return df_out.round(2)


# ----------------------------------------------------------------------------
# 论文表 3 / 紧急购电事件
# ----------------------------------------------------------------------------

def make_emergency_events(result_df, tol: float = EM_TOL_KW) -> pd.DataFrame:
    """
    把逐时段紧急购电功率合并为连续时间段事件。
    返回列：date, start_slot, end_slot, 时间段, 紧急购电量(kWh)
    """
    events = []
    for d, g in result_df.groupby("date", sort=True):
        g = g.sort_values("slot").reset_index(drop=True)
        active = (g["P_em_kW"].values > tol).astype(int)
        t = 0
        n = len(active)
        while t < n:
            if active[t] == 0:
                t += 1
                continue
            s = t
            while t + 1 < n and active[t + 1] == 1:
                t += 1
            e = t
            kwh = float((g.loc[s:e, "P_em_kW"] * HOURS_PER_SLOT).sum())
            events.append({
                "date": pd.Timestamp(d),
                "start_slot": s, "end_slot": e,
                "时间段": slot_range_label(s, e),
                "紧急购电量": kwh,
            })
            t += 1
    return pd.DataFrame(events)


def make_paper_table3_q2(result_df, key_dates=KEY_DATES,
                         tol: float = EM_TOL_KW) -> pd.DataFrame:
    """表 3：4 个指定日期的紧急购电时间段与购电量（长表，同日事件并列）。"""
    ev = make_emergency_events(result_df, tol=tol)
    rows = []
    for d in key_dates:
        sub = ev[ev["date"] == pd.Timestamp(d)]
        if sub.empty:
            rows.append({"日期": _date_label(d), "时间段": "无", "紧急购电量": 0.0})
        else:
            for i, (_, r) in enumerate(sub.iterrows()):
                rows.append({
                    "日期": _date_label(d) if i == 0 else "",
                    "时间段": r["时间段"],
                    "紧急购电量": round(float(r["紧急购电量"]), 2),
                })
    return pd.DataFrame(rows)


# ----------------------------------------------------------------------------
# result2.xlsx
# ----------------------------------------------------------------------------

def write_result2_xlsx(result_df, summary_df, out_path: Path = None,
                       template_path: Path = None):
    """复制模板并写入三个工作表。"""
    if template_path is None:
        template_path = RESULT_TEMPLATE_DIR / "result2.xlsx"
    if out_path is None:
        out_path = ROOT / "result2.xlsx"
    shutil.copy(template_path, out_path)
    wb = load_workbook(out_path)

    dates = sorted(result_df["date"].unique())

    # ---- 工作表 1：计划购电量（335 行 × 147 列） ----
    ws1 = wb["计划购电量"]
    for i, d in enumerate(dates):
        df = _day_df(result_df, d)
        row = i + 2
        ws1.cell(row=row, column=1, value=pd.Timestamp(d).to_pydatetime())
        kwh = df["P_plan_kW"].values * HOURS_PER_SLOT
        for t in range(SLOTS_PER_DAY):
            ws1.cell(row=row, column=t + 2, value=float(kwh[t]))
        ws1.cell(row=row, column=146, value=float(kwh.sum()))
        ws1.cell(row=row, column=147,
                 value=float((df["P_plan_kW"].values * df["price"].values
                              * HOURS_PER_SLOT).sum()))

    # ---- 工作表 2：充放电量（每天 6 个 4 小时时段 + 0:00/24:00 储电量） ----
    ws2 = wb["充放电量"]
    # 清空模板占位行
    if ws2.max_row > 1:
        ws2.delete_rows(2, ws2.max_row - 1)
    r = 2
    for d in dates:
        df = _day_df(result_df, d)
        summ = summary_df[summary_df["date"] == pd.Timestamp(d)].iloc[0]
        for j, (label, start, end) in enumerate(PAPER_TABLE2_RANGES):
            ws2.cell(row=r, column=1,
                     value=pd.Timestamp(d).to_pydatetime() if j == 0 else None)
            ws2.cell(row=r, column=2, value=label)
            ws2.cell(row=r, column=3, value=float(
                (df.loc[start:end, "P_ch_kW"] * HOURS_PER_SLOT).sum()))
            ws2.cell(row=r, column=4, value=float(
                (df.loc[start:end, "P_dis_kW"] * HOURS_PER_SLOT).sum()))
            if j == 0:
                ws2.cell(row=r, column=5, value=pd.Timestamp(d).to_pydatetime())
                ws2.cell(row=r, column=6, value=float(summ["e0"]))
            if j == 1:
                ws2.cell(row=r, column=5, value="24:00")
                ws2.cell(row=r, column=6, value=float(summ["e_end"]))
            r += 1

    # ---- 工作表 3：紧急购电量（日期 | 时间段 | 购电量） ----
    ws3 = wb["紧急购电量"]
    if ws3.max_row > 1:
        ws3.delete_rows(2, ws3.max_row - 1)
    ev = make_emergency_events(result_df)
    r = 2
    for _, e in ev.iterrows():
        ws3.cell(row=r, column=1, value=pd.Timestamp(e["date"]).to_pydatetime())
        ws3.cell(row=r, column=2, value=e["时间段"])
        ws3.cell(row=r, column=3, value=round(float(e["紧急购电量"]), 2))
        r += 1

    wb.save(out_path)
    wb.close()
    return out_path


# ----------------------------------------------------------------------------
# 一键输出
# ----------------------------------------------------------------------------

def format_all(result_df, summary_df, out_dir: Path = None):
    if out_dir is None:
        out_dir = ROOT
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    t1 = make_paper_table1_q2(result_df, summary_df)
    t2 = make_paper_table2_q2(result_df, summary_df)
    t3 = make_paper_table3_q2(result_df)
    events = make_emergency_events(result_df)

    xlsx = write_result2_xlsx(result_df, summary_df, out_path=out_dir / "result2.xlsx")
    t1.to_csv(out_dir / "table1_q2.csv", index=False, encoding="utf-8-sig")
    t2.to_csv(out_dir / "table2_q2.csv", index=False, encoding="utf-8-sig")
    t3.to_csv(out_dir / "table3_q2.csv", index=False, encoding="utf-8-sig")
    events.to_csv(out_dir / "emergency_events_q2.csv", index=False, encoding="utf-8-sig")
    return {"table1": t1, "table2": t2, "table3": t3,
            "events": events, "result_xlsx": xlsx}


def print_paper_tables(result_df, summary_df):
    t1 = make_paper_table1_q2(result_df, summary_df)
    t2 = make_paper_table2_q2(result_df, summary_df)
    t3 = make_paper_table3_q2(result_df)
    print("\n===== 表 1：指定日期的计划购电量及全天购电量/购电费（kWh / 元） =====")
    print(t1.to_string(index=False))
    print("\n===== 表 2：指定日期各时段充放电量及 0:00/24:00 储电量（kWh） =====")
    print(t2.to_string(index=False))
    print("\n===== 表 3：指定日期的紧急购电量（kWh） =====")
    print(t3.to_string(index=False))


if __name__ == "__main__":
    import pandas as pd
    rd = pd.read_csv(ROOT / "data" / "processed" / "result2_details.csv",
                     parse_dates=["date"])
    sd = pd.read_csv(ROOT / "data" / "processed" / "result2_summary.csv",
                     parse_dates=["date"])
    res = format_all(rd, sd)
    print_paper_tables(rd, sd)
    print(f"\n结果已保存到: {res['result_xlsx']}")
